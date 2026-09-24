"""Importação de turmas e alunos por planilha (CSV/XLSX).

Uma linha representa um aluno dentro de uma turma. O pipeline é sempre o
mesmo para prévia (dry-run) e confirmação: ler -> normalizar -> validar ->
resolver oficina/turma/aluno contra o banco. A única diferença é que a
confirmação, além disso, GRAVA — e só grava se não houver erro nem conflito.

Reaproveita as mesmas regras dos services de cadastro (organização sempre da
sessão, `get_for_org`, `normalize_email`), mas monta as entidades diretamente
(sem os commits por chamada de `catalog`/`students`), porque a importação
precisa ser atômica em lote: um único commit no fim, rollback de tudo em erro.
"""

import csv
import datetime as dt
import io
from dataclasses import dataclass, field

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import ClassGroup, Student, User, Workshop
from app.services.auth import normalize_email
from app.services.errors import InvalidData

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
STUDENT_STATUSES = ("ativo", "desistente")

REQUIRED_COLUMNS = ("oficina", "turma", "aluno_nome", "data_matricula")
OPTIONAL_COLUMNS = ("horario", "professor_email", "aluno_status")
TEMPLATE_COLUMNS = ("oficina", "turma", "horario", "professor_email", "aluno_nome", "data_matricula", "aluno_status")
TEMPLATE_EXAMPLE_ROWS = (
    ("Ballet", "Ballet 14h", "14:00", "larissa@example.com", "Ana Silva", "2026-09-01", "ativo"),
    ("Ballet", "Ballet 14h", "14:00", "larissa@example.com", "Beatriz Souza", "2026-09-01", "ativo"),
    ("Ballet", "Ballet 15h", "15:00", "larissa@example.com", "Carla Lima", "2026-09-01", "ativo"),
)

_MAX_LENGTHS = {"oficina": 120, "turma": 120, "horario": 120, "aluno_nome": 160}

# Sinônimos de cabeçalho aceitos (comparação exata após lower()+strip(), sem fuzzy
# matching): planilhas de organizações diferentes nomeiam as mesmas colunas de jeitos
# óbvios e não ambíguos. Só o rótulo muda — o formato do valor esperado na coluna
# continua o mesmo (ex.: "professora" ainda precisa conter um e-mail, não um nome).
# "oficineiro_email" é o nome antigo da coluna (formato oficial anterior a esta etapa
# de unificação de terminologia); mantido como alias legado para não quebrar planilhas
# já em uso — nunca remover sem aviso.
HEADER_ALIASES = {
    "aluno": "aluno_nome",
    "nome": "aluno_nome",
    "nome do aluno": "aluno_nome",
    "professor": "professor_email",
    "professora": "professor_email",
    "email professor": "professor_email",
    "e-mail do professor": "professor_email",
    "oficineiro_email": "professor_email",
}


def _norm(value) -> str:
    return "" if value is None else str(value).strip()


def _norm_key(value: str) -> str:
    """Chave de comparação tolerante a maiúsculas/espaços (não é o valor salvo)."""
    return " ".join(value.split()).casefold()


def _parse_date(raw) -> dt.date:
    if isinstance(raw, dt.datetime):
        return raw.date()
    if isinstance(raw, dt.date):
        return raw
    text = _norm(raw)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


# ------------------------------------------------------------------- leitura
def _read_csv(raw: bytes) -> tuple[list[str], list[list]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InvalidData("Não foi possível ler o arquivo CSV. Salve-o como CSV UTF-8 e tente novamente.") from exc
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise InvalidData("A planilha está vazia.")
    header = [_norm(h).lower() for h in rows[0]]
    return header, rows[1:]


def _read_xlsx(raw: bytes) -> tuple[list[str], list[list]]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        data_rows = [list(r) for r in rows_iter]
    except Exception as exc:
        raise InvalidData("Não foi possível ler o arquivo .xlsx. Verifique se ele não está corrompido.") from exc
    if header_row is None:
        raise InvalidData("A planilha está vazia.")
    header = [_norm(h).lower() for h in header_row]
    return header, data_rows


def _read_rows(filename: str, raw: bytes) -> tuple[list[str], list[list]]:
    if len(raw) > MAX_FILE_BYTES:
        raise InvalidData(f"Arquivo muito grande (máximo {MAX_FILE_BYTES // (1024 * 1024)} MB).")
    lower = filename.lower()
    if lower.endswith(".csv"):
        header, rows = _read_csv(raw)
    elif lower.endswith(".xlsx"):
        header, rows = _read_xlsx(raw)
    else:
        raise InvalidData("Envie um arquivo .csv ou .xlsx.")
    header = [HEADER_ALIASES.get(h, h) for h in header]
    seen_cols = set()
    for h in header:
        if h in seen_cols:
            raise InvalidData(f'Coluna "{h}" aparece duplicada no cabeçalho (direto ou por sinônimo).')
        seen_cols.add(h)
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise InvalidData(f"Cabeçalho incompleto: faltam as colunas {', '.join(missing)}.")
    rows = [r for r in rows if any(_norm(c) for c in r)]  # ignora linhas totalmente em branco
    if not rows:
        raise InvalidData("A planilha não tem nenhuma linha de dados.")
    if len(rows) > MAX_ROWS:
        raise InvalidData(f"Muitas linhas (máximo {MAX_ROWS}).")
    return header, rows


# --------------------------------------------------------------- estruturas
@dataclass
class RowResult:
    row: int
    oficina: str
    turma: str
    aluno: str
    status: str  # criar | existente | conflito | erro
    message: str = ""

    def to_dict(self) -> dict:
        return {"row": self.row, "oficina": self.oficina, "turma": self.turma, "aluno": self.aluno,
                "status": self.status, "message": self.message}


@dataclass
class _WorkshopPlan:
    action: str  # "existing" | "create"
    name: str
    instructor_id: int | None
    obj: Workshop | None = None


@dataclass
class _ClassPlan:
    action: str
    workshop_key: str
    name: str
    schedule: str
    obj: ClassGroup | None = None


@dataclass
class ImportResult:
    filename: str
    total_rows: int
    rows: list[RowResult] = field(default_factory=list)
    workshops_to_create: int = 0
    classes_to_create: int = 0
    students_to_create: int = 0
    students_existing: int = 0
    students_conflict: int = 0
    errors: int = 0
    imported: bool = False

    @property
    def ok_to_import(self) -> bool:
        return self.errors == 0 and self.students_conflict == 0

    def to_dict(self) -> dict:
        return {
            "filename": self.filename,
            "total_rows": self.total_rows,
            "ok_to_import": self.ok_to_import,
            "imported": self.imported,
            "summary": {
                "workshops_to_create": self.workshops_to_create,
                "classes_to_create": self.classes_to_create,
                "students_to_create": self.students_to_create,
                "students_existing": self.students_existing,
                "students_conflict": self.students_conflict,
                "errors": self.errors,
            },
            "rows": [r.to_dict() for r in self.rows],
        }


# -------------------------------------------------------------------- linha
def _resolve_instructor(db: Session, email: str, organization_id: int) -> int | None:
    """None só quando a coluna veio vazia; e-mail preenchido tem que resolver ou vira erro."""
    normalized = normalize_email(email)
    user = db.scalar(select(User).where(User.email == normalized))
    # Igual a outra organização: mesma mensagem de "não encontrado" (nunca revela o outro tenant).
    if user is None or user.organization_id != organization_id or user.instructor_id is None:
        raise InvalidData("Professor não encontrado. Cadastre o professor antes de importar.")
    return user.instructor_id


def _process_row(
    db: Session, organization_id: int, row_num: int, values: dict[str, str],
    workshops: dict[str, _WorkshopPlan], classes: dict[str, _ClassPlan], seen: dict[tuple, int],
    resolved: dict[int, dict],
) -> RowResult:
    oficina, turma, aluno = values["oficina"], values["turma"], values["aluno_nome"]
    result = RowResult(row=row_num, oficina=oficina, turma=turma, aluno=aluno, status="erro")

    missing = [c for c in REQUIRED_COLUMNS if not values[c]]
    if missing:
        result.message = f"Faltando: {', '.join(missing)}."
        return result

    for field_name, limit in _MAX_LENGTHS.items():
        if len(values[field_name]) > limit:
            result.message = f"{field_name} muito longo (máximo {limit} caracteres)."
            return result

    try:
        enrollment_date = _parse_date(values["data_matricula"])
    except ValueError:
        result.message = f"Data de matrícula inválida: '{values['data_matricula']}'."
        return result

    status = values["aluno_status"] or "ativo"
    if status not in STUDENT_STATUSES:
        result.message = f"Situação inválida: '{status}' (use ativo ou desistente)."
        return result

    dup_key = (_norm_key(oficina), _norm_key(turma), _norm_key(aluno))
    if dup_key in seen:
        result.message = f"Aluno duplicado nesta planilha (já aparece na linha {seen[dup_key]})."
        return result
    seen[dup_key] = row_num

    instructor_id = None
    if values["professor_email"]:
        try:
            instructor_id = _resolve_instructor(db, values["professor_email"], organization_id)
        except InvalidData as exc:
            result.message = exc.message
            return result

    workshop_key = _norm_key(oficina)
    plan = workshops.get(workshop_key)
    if plan is None:
        existing = db.scalar(
            select(Workshop).where(Workshop.organization_id == organization_id, Workshop.name.ilike(oficina))
        )
        plan = (_WorkshopPlan("existing", oficina, existing.instructor_id, existing) if existing
                else _WorkshopPlan("create", oficina, instructor_id))
        workshops[workshop_key] = plan

    class_key = (workshop_key, _norm_key(turma))
    cplan = classes.get(class_key)
    if cplan is None:
        if plan.action == "existing":
            existing_class = db.scalar(
                select(ClassGroup).where(
                    ClassGroup.organization_id == organization_id,
                    ClassGroup.workshop_id == plan.obj.id,
                    ClassGroup.name.ilike(turma),
                )
            )
        else:
            existing_class = None
        cplan = (_ClassPlan("existing", workshop_key, turma, values["horario"], existing_class) if existing_class
                 else _ClassPlan("create", workshop_key, turma, values["horario"]))
        classes[class_key] = cplan

    # aluno já existe? na turma-alvo = já existente (ignora); em outra turma = conflito (ação manual)
    same_name = list(db.scalars(
        select(Student).where(Student.organization_id == organization_id, Student.name.ilike(aluno))
    ))
    if cplan.action == "existing":
        in_target = next((s for s in same_name if s.class_id == cplan.obj.id), None)
        if in_target:
            result.status, result.message = "existente", "Aluno já matriculado nesta turma."
            return result
    elsewhere = same_name[0] if same_name else None
    if elsewhere:
        result.status = "conflito"
        result.message = f"Já existe um(a) aluno(a) com este nome na turma \"{elsewhere.class_group.name}\"."
        return result

    result.status = "criar"
    resolved[row_num] = {"class_key": class_key, "enrollment_date": enrollment_date}
    return result


# --------------------------------------------------------------- pipeline
def run_import(db: Session, filename: str, raw: bytes, organization_id: int, *, dry_run: bool) -> ImportResult:
    """Sempre lê, normaliza, valida e resolve. Só grava se dry_run=False e não houver
    erro/conflito — nunca importação parcial: um commit no final, rollback em qualquer falha."""
    header, data_rows = _read_rows(filename, raw)
    result = ImportResult(filename=filename, total_rows=len(data_rows))
    workshops: dict[str, _WorkshopPlan] = {}
    classes: dict[tuple, _ClassPlan] = {}
    seen: dict[tuple, int] = {}
    resolved: dict[int, dict] = {}  # linhas com status "criar": class_key + enrollment_date já validados

    for i, raw_row in enumerate(data_rows):
        row_num = i + 2  # linha 1 = cabeçalho
        values = {}
        for col in REQUIRED_COLUMNS + OPTIONAL_COLUMNS:
            idx = header.index(col) if col in header else -1
            cell = raw_row[idx] if 0 <= idx < len(raw_row) else None
            # data_matricula como célula de data real do Excel não pode virar string aqui:
            # perderia o tipo antes de chegar em _parse_date (datetime vira "2026-09-22 00:00:00",
            # que não bate com nenhum formato aceito). Só normaliza texto (CSV, ou célula vazia).
            values[col] = cell if col == "data_matricula" and isinstance(cell, dt.date) else _norm(cell)
        row_result = _process_row(db, organization_id, row_num, values, workshops, classes, seen, resolved)
        result.rows.append(row_result)
        if row_result.status == "erro":
            result.errors += 1
        elif row_result.status == "conflito":
            result.students_conflict += 1
        elif row_result.status == "existente":
            result.students_existing += 1

    result.workshops_to_create = sum(1 for p in workshops.values() if p.action == "create")
    result.classes_to_create = sum(1 for p in classes.values() if p.action == "create")
    result.students_to_create = sum(1 for r in result.rows if r.status == "criar")

    if dry_run or not result.ok_to_import:
        return result

    try:
        for plan in workshops.values():
            if plan.action == "create":
                plan.obj = Workshop(
                    organization_id=organization_id, name=plan.name,
                    instructor_id=plan.instructor_id, active=True,
                )
                db.add(plan.obj)
        db.flush()

        for cplan in classes.values():
            if cplan.action == "create":
                cplan.obj = ClassGroup(
                    organization_id=organization_id, workshop_id=workshops[cplan.workshop_key].obj.id,
                    name=cplan.name, schedule=cplan.schedule or None, active=True,
                )
                db.add(cplan.obj)
        db.flush()

        for row_result in result.rows:
            if row_result.status != "criar":
                continue
            info = resolved[row_result.row]
            cplan = classes[info["class_key"]]
            db.add(Student(
                organization_id=organization_id, class_id=cplan.obj.id, name=row_result.aluno,
                enrollment_date=info["enrollment_date"], status="ativo",
            ))
        db.commit()
    except Exception:
        db.rollback()
        raise

    result.imported = True
    return result


def build_template(fmt: str) -> bytes:
    """CSV ou XLSX com o cabeçalho oficial e linhas de exemplo."""
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(TEMPLATE_COLUMNS)
        writer.writerows(TEMPLATE_EXAMPLE_ROWS)
        return buf.getvalue().encode("utf-8-sig")
    if fmt == "xlsx":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Importação"
        sheet.append(TEMPLATE_COLUMNS)
        for row in TEMPLATE_EXAMPLE_ROWS:
            sheet.append(row)
        out = io.BytesIO()
        workbook.save(out)
        return out.getvalue()
    raise InvalidData("Formato de modelo inválido.")
