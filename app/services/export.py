"""Exportação DOCX do relatório, seguindo a estrutura do modelo oficial."""

import datetime as dt
import io

from docx import Document
from docx.shared import Pt

from app.services.periods import MONTH_NAMES
from app.services.reports import INFRASTRUCTURE_LABELS, MATERIALS_LABELS, METHODOLOGY_LABELS


def _box(checked: bool) -> str:
    return "☒" if checked else "☐"


def as_date(value) -> dt.date | None:
    if value is None or isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value)[:10])


def _label_line(doc, label: str, value: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{label} ").bold = True
    p.add_run(value or "")


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def _text_block(doc, text: str) -> None:
    for line in (text or "").splitlines() or [""]:
        doc.add_paragraph(line)


def build_docx(report: dict) -> bytes:
    ident, fields, part = report["identification"], report["fields"], report["participation"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title_run = title.add_run("Relatório Mensal de Atividades: Oficinas Itinerantes")
    title_run.bold = True
    title_run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.add_run("Mês de Referência: ").bold = True
    p.add_run("  ".join(f"{_box(i + 1 == report['month'])} {m}" for i, m in enumerate(MONTH_NAMES)))

    _heading(doc, "1. Identificação:")
    _label_line(doc, "Professor(a):", ident["instructor_name"])
    _label_line(doc, "Oficina/Linguagem:", " / ".join(x for x in (ident["workshop_name"], ident["language"]) if x))
    _label_line(doc, "Local/Comunidade atendida:", ident["location"])

    _heading(doc, "2. Resumo das Atividades Realizadas")
    _text_block(doc, fields["activities_summary"])

    _heading(doc, "3. Dados de Participação")
    _label_line(doc, "Nº de alunos matriculados:", part["enrolled"])
    _label_line(doc, "Média de presença no mês:", part["average"])
    _label_line(doc, "Novas matrículas/Desistências:", part["changes"])

    _heading(doc, "4. Metodologia Aplicada e Resultados Observados")
    for key, (label, desc) in METHODOLOGY_LABELS.items():
        p = doc.add_paragraph()
        p.add_run(f"{_box(fields['methodology_rating'] == key)} ")
        p.add_run(f"{label}: ").bold = True
        p.add_run(desc)
    p = doc.add_paragraph()
    p.add_run("Observações:").bold = True
    _text_block(doc, fields["methodology_notes"])

    _heading(doc, "5. Desafios e Logística")
    _label_line(
        doc,
        "Infraestrutura:",
        "  ".join(f"({'X' if fields['infrastructure'] == k else ' '}) {v}" for k, v in INFRASTRUCTURE_LABELS.items()),
    )
    _label_line(
        doc,
        "Materiais:",
        "  ".join(f"({'X' if fields['materials'] == k else ' '}) {v}" for k, v in MATERIALS_LABELS.items()),
    )
    _label_line(doc, "Comentários:", "")
    _text_block(doc, fields["logistics_comments"])

    _heading(doc, "6. Destaques e Ocorrências")
    _text_block(doc, fields["highlights_occurrences"])

    _heading(doc, "7. Planejamento para o Próximo Mês")
    _text_block(doc, fields["next_month_planning"])

    report_date = as_date(fields["report_date"])
    p = doc.add_paragraph()
    p.add_run("Data: ").bold = True
    p.add_run(report_date.strftime("%d/%m/%Y") if report_date else "____/____/______")
    p = doc.add_paragraph()
    p.add_run("Assinatura do Professor: ").bold = True
    p.add_run(f"{fields['signature_name']}  ______________________________")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
